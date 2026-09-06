"""构建 DOCX 报告 - 从 Markdown 转换"""

import re
import shutil
from pathlib import Path
from typing import Dict, Any, List, Optional

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn, nsmap
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# 字体配置
FONT_CHINESE = "微软雅黑"
FONT_ENGLISH = "Calibri"


def is_chinese_char(char: str) -> bool:
    """判断字符是否为中文"""
    return '\u4e00' <= char <= '\u9fff'


def set_run_font(run, text: str, size_pt: int = 11, bold: bool = False, italic: bool = False):
    """设置 run 的字体，中文用微软雅黑，英文用 Calibri"""
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    
    # 设置西文字体
    run.font.name = FONT_ENGLISH
    # 设置中文字体
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CHINESE)


def add_styled_paragraph(doc, text: str, size_pt: int = 11, bold: bool = False, 
                         italic: bool = False, alignment=None, space_after: int = 6):
    """添加带样式的段落"""
    para = doc.add_paragraph()
    if alignment:
        para.alignment = alignment
    para.paragraph_format.space_after = Pt(space_after)
    
    run = para.add_run()
    set_run_font(run, text, size_pt, bold, italic)
    return para


def add_cover_page(doc, title: str, engine_model: str, project_code: str, created_at: str):
    """添加封面页"""
    # 添加空行使内容垂直居中
    for _ in range(8):
        doc.add_paragraph()
    
    # 主标题
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run()
    set_run_font(title_run, title, size_pt=28, bold=True)
    
    # 空行
    for _ in range(3):
        doc.add_paragraph()
    
    # 发动机型号
    model_para = doc.add_paragraph()
    model_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    model_run = model_para.add_run()
    set_run_font(model_run, f"发动机型号：{engine_model}", size_pt=16)
    
    # 项目代码
    code_para = doc.add_paragraph()
    code_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    code_run = code_para.add_run()
    set_run_font(code_run, f"项目代码：{project_code}", size_pt=16)
    
    # 生成时间
    time_para = doc.add_paragraph()
    time_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    time_run = time_para.add_run()
    set_run_font(time_run, f"生成时间：{created_at}", size_pt=16)
    
    # 添加分页符
    doc.add_page_break()


def add_table_of_contents(doc):
    """添加目录页"""
    # 目录标题
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_run = toc_title.add_run()
    set_run_font(toc_run, "目  录", size_pt=22, bold=True)
    
    doc.add_paragraph()
    
    # 添加 TOC 域代码
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    # 占位文本
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    
    # 添加提示文本（目录将在打开时自动更新）
    hint_para = doc.add_paragraph()
    hint_run = hint_para.add_run()
    set_run_font(hint_run, '（目录将在打开文档时自动更新，如未显示请按 Ctrl+A 后按 F9）', size_pt=10, italic=True)
    
    run._r.append(fldChar3)
    
    # 设置文档属性：打开时更新域
    settings = doc.settings.element
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    settings.append(update_fields)
    
    # 添加分页符
    doc.add_page_break()


def build_docx_report(
    md_content: str,
    docx_path: str,
    assets_dir: str,
    engine_model: str = "",
    project_code: str = "",
    created_at: str = ""
) -> bool:
    """
    从 Markdown 内容构建 DOCX 报告
    
    Args:
        md_content: Markdown 报告内容
        docx_path: 输出 DOCX 文件路径
        assets_dir: 图片资源目录
        
    Returns:
        是否成功
    """
    if not HAS_DOCX:
        print("警告: python-docx 未安装，无法生成 DOCX 报告")
        return False
    
    try:
        doc = Document()
        assets_path = Path(assets_dir)
        
        # 从 Markdown 中提取元信息（如果未提供）
        if not engine_model or not project_code or not created_at:
            engine_model, project_code, created_at = extract_metadata_from_md(md_content)
        
        # 1. 添加封面页
        add_cover_page(doc, "涡喷发动机仿真分析报告", engine_model, project_code, created_at)
        
        # 2. 添加目录页
        add_table_of_contents(doc)
        
        # 3. 设置默认样式字体
        setup_document_styles(doc)
        
        # 4. 解析 Markdown 并转换为 DOCX
        lines = md_content.split('\n')
        i = 0
        in_table = False
        table_rows = []
        skip_header_section = True  # 跳过 Markdown 开头的元信息
        
        while i < len(lines):
            line = lines[i]
            
            # 跳过 Markdown 开头的元信息部分（到第一个 --- 为止）
            if skip_header_section:
                if line.startswith('---') and i > 0:
                    skip_header_section = False
                i += 1
                continue
            
            # 跳过第一个 # 标题（已在封面显示）
            if line.startswith('# ') and '仿真分析报告' in line:
                i += 1
                continue
            
            # 标题
            if line.startswith('## '):
                heading = doc.add_heading(line[3:].strip(), level=1)
                set_heading_font(heading, 1)
            elif line.startswith('### '):
                heading = doc.add_heading(line[4:].strip(), level=2)
                set_heading_font(heading, 2)
            elif line.startswith('#### '):
                heading = doc.add_heading(line[5:].strip(), level=3)
                set_heading_font(heading, 3)
            
            # 图片
            elif line.startswith('!['):
                match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
                if match:
                    alt_text, img_path = match.groups()
                    full_img_path = assets_path.parent / img_path
                    if full_img_path.exists():
                        try:
                            doc.add_picture(str(full_img_path), width=Inches(5.5))
                            # 添加图片说明
                            caption = doc.add_paragraph()
                            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            cap_run = caption.add_run()
                            set_run_font(cap_run, alt_text, size_pt=10, italic=True)
                        except Exception:
                            add_styled_paragraph(doc, f"[图片: {alt_text}]")
                    else:
                        add_styled_paragraph(doc, f"[图片: {alt_text}]")
            
            # 表格
            elif line.startswith('|'):
                if not in_table:
                    in_table = True
                    table_rows = []
                table_rows.append(line)
            elif in_table and not line.startswith('|'):
                if table_rows:
                    add_styled_table(doc, table_rows)
                in_table = False
                table_rows = []
                if line.strip() and not line.startswith('---'):
                    add_formatted_paragraph(doc, line)
            
            # 分隔线 - 跳过不添加
            elif line.startswith('---'):
                pass
            
            # 斜体注释（页脚）
            elif line.startswith('*') and line.endswith('*') and len(line) > 2:
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                set_run_font(run, line[1:-1], size_pt=10, italic=True)
            
            # 普通段落
            elif line.strip():
                add_formatted_paragraph(doc, line)
            
            i += 1
        
        # 处理最后的表格
        if in_table and table_rows:
            add_styled_table(doc, table_rows)
        
        # 保存文件
        doc.save(docx_path)
        return True
        
    except Exception as e:
        print(f"生成 DOCX 失败: {e}")
        import traceback
        traceback.print_exc()
        return False


def extract_metadata_from_md(md_content: str) -> tuple:
    """从 Markdown 内容中提取元信息"""
    engine_model = ""
    project_code = ""
    created_at = ""
    
    for line in md_content.split('\n')[:20]:
        if '发动机型号' in line:
            match = re.search(r'[：:]\s*(.+)', line)
            if match:
                engine_model = match.group(1).strip()
        elif '项目代码' in line:
            match = re.search(r'[：:]\s*(.+)', line)
            if match:
                project_code = match.group(1).strip()
        elif '生成时间' in line:
            match = re.search(r'[：:]\s*(.+)', line)
            if match:
                created_at = match.group(1).strip()
    
    return engine_model, project_code, created_at


def setup_document_styles(doc):
    """设置文档默认样式"""
    # 设置正文样式
    style = doc.styles['Normal']
    font = style.font
    font.name = FONT_ENGLISH
    font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CHINESE)


def set_heading_font(heading, level: int):
    """设置标题字体"""
    sizes = {1: 16, 2: 14, 3: 12}
    size = sizes.get(level, 12)
    
    for run in heading.runs:
        run.font.name = FONT_ENGLISH
        run.font.size = Pt(size)
        run.font.bold = True
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CHINESE)


def add_formatted_paragraph(doc, text: str):
    """添加带格式的段落"""
    para = doc.add_paragraph()
    
    # 处理加粗和斜体
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run()
            set_run_font(run, part[2:-2], bold=True)
        elif part.startswith('*') and part.endswith('*'):
            run = para.add_run()
            set_run_font(run, part[1:-1], italic=True)
        else:
            run = para.add_run()
            set_run_font(run, part)


def add_styled_table(doc, rows: List[str]):
    """添加带样式的表格"""
    if len(rows) < 2:
        return
    
    data = []
    for row in rows:
        if '---' in row:
            continue
        cells = [c.strip() for c in row.split('|')[1:-1]]
        if cells:
            data.append(cells)
    
    if not data:
        return
    
    num_cols = len(data[0])
    table = doc.add_table(rows=len(data), cols=num_cols)
    table.style = 'Table Grid'
    
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < len(row.cells):
                cell = row.cells[j]
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run()
                clean_text = re.sub(r'\*\*([^*]+)\*\*', r'\1', cell_text)
                clean_text = re.sub(r'\*([^*]+)\*', r'\1', clean_text)
                set_run_font(run, clean_text, size_pt=10, bold=(i == 0))


def process_inline_formatting(text: str) -> str:
    """处理行内格式（移除 Markdown 标记）"""
    # 移除加粗
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    # 移除斜体
    text = re.sub(r'\*([^*]+)\*', r'\1', text)
    # 移除代码
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text


def add_table_from_markdown(doc, rows: List[str]):
    """从 Markdown 表格行添加 DOCX 表格"""
    if len(rows) < 2:
        return
    
    # 解析表格数据
    data = []
    for row in rows:
        if '---' in row:
            continue  # 跳过分隔行
        cells = [c.strip() for c in row.split('|')[1:-1]]
        if cells:
            data.append(cells)
    
    if not data:
        return
    
    # 创建表格
    num_cols = len(data[0])
    table = doc.add_table(rows=len(data), cols=num_cols)
    table.style = 'Table Grid'
    
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            if j < len(row.cells):
                row.cells[j].text = process_inline_formatting(cell_text)


def build_replacements(facts: Dict[str, Any], filled_slots: Dict[str, str]) -> Dict[str, str]:
    """构建文本替换映射"""
    project = facts.get("project", {})
    manifest = facts.get("manifest", {})
    
    replacements = {
        # 基本信息
        "[---涡喷发动机型号---]": project.get("engine_model_name", "未指定型号"),
        "[---分析目的1---]": filled_slots.get("analysis_purpose_1", "对涡喷发动机内部流场进行数值模拟分析"),
        "[---分析目的2---]": filled_slots.get("analysis_purpose_2", "验证设计方案的合理性"),
        
        # 知识库查询结果
        "【知识库中的某个仿真规范】": "常用CFD仿真规范",
        "【知识库中查询空气】": "标准大气条件",
        "【知识库中的查询的评定准则】": filled_slots.get("evaluation_criteria_1", ""),
        
        # 结果数据
        "【分析结果数据】": generate_result_summary(facts),
        
        # 结论
        "【结合知识库和结果数据生成的结论】": filled_slots.get("conclusion", ""),
        
        # 暂不填写
        "【暂时不加】": "本模板暂不填写",
    }
    
    return replacements


def replace_in_paragraph(
    para,
    replacements: Dict[str, str],
    images: Dict[str, List[str]],
    image_mapping: Dict[str, str],
    out_path: Path
):
    """在段落中进行替换"""
    text = para.text
    
    # 检查是否是图片占位符
    if text.strip() == "【图片】":
        # 尝试插入结果图片
        insert_image_in_paragraph(para, images, image_mapping, out_path, "variables")
        return
    
    if text.strip() == "【整体网格图片】":
        insert_image_in_paragraph(para, images, image_mapping, out_path, "mesh")
        return
    
    if text.strip() == "【以图片分别陈列所有模块】":
        insert_image_in_paragraph(para, images, image_mapping, out_path, "geometry")
        return
    
    # 文本替换
    for pattern, replacement in replacements.items():
        if pattern in text:
            # 需要在 runs 中查找并替换
            replace_text_in_runs(para, pattern, replacement)


def replace_text_in_runs(para, pattern: str, replacement: str):
    """在段落的 runs 中替换文本"""
    # 简单方法：直接替换整个段落文本
    full_text = para.text
    if pattern in full_text:
        new_text = full_text.replace(pattern, replacement)
        
        # 清除所有 runs
        for run in para.runs:
            run.text = ""
        
        # 在第一个 run 中设置新文本（保留格式）
        if para.runs:
            para.runs[0].text = new_text
        else:
            para.add_run(new_text)


def insert_image_in_paragraph(
    para,
    images: Dict[str, List[str]],
    image_mapping: Dict[str, str],
    out_path: Path,
    category: str
):
    """在段落位置插入图片"""
    img_list = images.get(category, [])
    
    if not img_list:
        # 没有图片，替换为提示文本
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = f"（数据包未包含{category}图片）"
        return
    
    # 清除原有文本
    for run in para.runs:
        run.text = ""
    
    # 插入第一张图片
    img_path = img_list[0]
    
    # 检查是否在 assets 目录
    mapped_path = image_mapping.get(img_path)
    if mapped_path:
        img_path = out_path / mapped_path
    else:
        img_path = Path(img_path)
    
    if img_path.exists():
        try:
            run = para.add_run()
            run.add_picture(str(img_path), width=Inches(5.5))
        except Exception:
            para.add_run(f"（图片加载失败: {img_path.name}）")
    else:
        para.add_run(f"（图片文件不存在）")


def generate_result_summary(facts: Dict[str, Any]) -> str:
    """生成结果数据摘要"""
    variables = facts.get("variables", [])
    
    if not variables:
        return "数据包未包含变量数据"
    
    lines = []
    for var in variables:
        name = var.get("name", "")
        guess = var.get("guess_variableName", "")
        range_min = var.get("range_min")
        range_max = var.get("range_max")
        
        if range_min is not None and range_max is not None:
            lines.append(f"{name}({guess}): {range_min:.4f} ~ {range_max:.4f}")
        else:
            lines.append(f"{name}({guess}): 范围未知")
    
    return "；".join(lines)


def update_air_properties_table(doc):
    """更新空气物性表（表3）"""
    # 查找包含"物理属性"或"材料"的表格
    for table in doc.tables:
        # 检查表格是否是材料属性表
        first_cell_text = ""
        if table.rows and table.rows[0].cells:
            first_cell_text = table.rows[0].cells[0].text.strip()
        
        # 如果找到类似材料表的结构，尝试更新
        if "材料" in first_cell_text or "属性" in first_cell_text or "参数" in first_cell_text:
            try:
                update_table_with_air_properties(table)
            except Exception:
                pass


def update_table_with_air_properties(table):
    """用空气物性数据更新表格"""
    # 获取表格行数
    row_count = len(table.rows)
    
    # 空气物性数据
    properties = [
        ("温度", "288.15", "K"),
        ("密度", "1.225", "kg/m³"),
        ("动力粘度", "1.789×10⁻⁵", "Pa·s"),
        ("定压比热", "1005", "J/(kg·K)"),
        ("导热系数", "0.0257", "W/(m·K)"),
        ("比热比 γ", "1.4", "-"),
        ("气体常数 R", "287.05", "J/(kg·K)")
    ]
    
    # 尝试填充数据（跳过表头）
    for i, (name, value, unit) in enumerate(properties):
        row_idx = i + 1  # 跳过表头
        if row_idx < row_count:
            row = table.rows[row_idx]
            if len(row.cells) >= 3:
                row.cells[0].text = name
                row.cells[1].text = value
                row.cells[2].text = unit


def build_docx_from_markdown(
    md_path: str,
    template_path: str,
    out_dir: str
) -> Dict[str, Any]:
    """
    使用 pandoc 从 Markdown 生成 DOCX（备选方案）
    
    Args:
        md_path: Markdown 文件路径
        template_path: 模板文件路径（作为样式参考）
        out_dir: 输出目录
        
    Returns:
        包含 docx_path, error 的字典
    """
    import subprocess
    
    out_path = Path(out_dir)
    docx_file = out_path / "report.docx"
    
    cmd = [
        "pandoc",
        md_path,
        "-o", str(docx_file),
        "--reference-doc", template_path
    ]
    
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if result.returncode == 0:
            return {
                "docx_path": str(docx_file),
                "error": None
            }
        else:
            return {
                "docx_path": None,
                "error": f"pandoc 错误: {result.stderr}"
            }
    except FileNotFoundError:
        return {
            "docx_path": None,
            "error": "pandoc 未安装"
        }
    except subprocess.TimeoutExpired:
        return {
            "docx_path": None,
            "error": "pandoc 执行超时"
        }
    except Exception as e:
        return {
            "docx_path": None,
            "error": str(e)
        }
